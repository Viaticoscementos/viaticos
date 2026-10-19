"""Genera el memorandum y la tabla de liquidacion en el formato real de CENOSA
(basado en el memo de Mario Urquia, Ejecutivo de Ventas Senior, centro de
costo V60140).

Salidas:
  - memorandum (.xlsx) por solicitud o por lote aprobado
  - tabla de liquidacion (.xlsx) con el detalle de facturas
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

from core.config import REPORTES_DIR, NOMBRE_EMPRESA_DEFAULT
from core import db

_AZUL = "1F4E79"
_GRIS = "D9E1F2"
_borde = Border(*(Side(style="thin", color="999999"),) * 4)


def _meses():
    return ["", "enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
            "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def _fecha_larga(d=None):
    d = d or datetime.now()
    return f"San Pedro Sula, {d.day} de {_meses()[d.month]} de {d.year}"


def _l(v):
    try:
        return f"L {float(v):,.2f}"
    except (TypeError, ValueError):
        return "L 0.00"


def generar_memo(solicitudes, empleado=None, centro_costo="", puesto="",
                 ruta_salida=None):
    """Memorandum de reembolso de viaticos en el formato de la empresa."""
    if isinstance(solicitudes, dict):
        solicitudes = [solicitudes]
    emp = empleado or (solicitudes[0].get("empleado") if solicitudes else "")
    datos_emp = db.buscar_empleado(emp) or {}
    puesto = puesto or datos_emp.get("puesto", "")
    centro_costo = centro_costo or datos_emp.get("departamento", "")

    wb = Workbook()
    ws = wb.active
    ws.title = "Memorandum"
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEF", (4, 22, 20, 16, 16, 16)):
        ws.column_dimensions[col].width = w

    tit = Font(bold=True, size=13, color=_AZUL)
    neg = Font(bold=True)

    ws.merge_cells("B1:F1")
    ws["B1"] = NOMBRE_EMPRESA_DEFAULT
    ws["B1"].font = tit
    ws.merge_cells("B2:F2")
    ws["B2"] = "MEMORANDUM - LIQUIDACION DE VIATICOS"
    ws["B2"].font = Font(bold=True, size=11)
    ws.merge_cells("B3:F3")
    ws["B3"] = _fecha_larga()
    ws["B3"].font = Font(italic=True, size=9)

    fila = 5
    enc = [("Para:", "Departamento de Contabilidad"),
           ("De:", emp),
           ("Puesto:", puesto or "-"),
           ("Centro de costo:", centro_costo or "-"),
           ("Asunto:", "Reembolso / liquidacion de gastos de viaticos")]
    for et, val in enc:
        ws[f"B{fila}"] = et
        ws[f"B{fila}"].font = neg
        ws.merge_cells(f"C{fila}:F{fila}")
        ws[f"C{fila}"] = val
        fila += 1

    fila += 1
    cols = ["Folio", "Fecha", "Categoria", "Negocio / proveedor",
            "Monto factura", "Reembolsable"]
    for i, c in enumerate(cols):
        cel = ws.cell(row=fila, column=2 + i, value=c)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=_AZUL)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = _borde
    fila += 1

    total_f = total_r = 0.0
    for idx, s in enumerate(solicitudes):
        valores = [s.get("folio", ""), s.get("fecha_compra", "") or "",
                   (s.get("categoria") or "").title(), s.get("negocio") or "-",
                   float(s.get("monto") or 0), float(s.get("monto_reembolsable") or 0)]
        total_f += valores[4]
        total_r += valores[5]
        for i, v in enumerate(valores):
            cel = ws.cell(row=fila, column=2 + i, value=v)
            cel.border = _borde
            if idx % 2:
                cel.fill = PatternFill("solid", fgColor=_GRIS)
            if i >= 4:
                cel.number_format = '"L" #,##0.00'
                cel.alignment = Alignment(horizontal="right")
        fila += 1

    ws.cell(row=fila, column=2, value="TOTAL").font = neg
    ws.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
    tf = ws.cell(row=fila, column=5, value=total_f)
    tf.number_format = '"L" #,##0.00'
    tf.font = neg
    tr = ws.cell(row=fila, column=6, value=total_r)
    tr.number_format = '"L" #,##0.00'
    tr.font = neg
    tr.fill = PatternFill("solid", fgColor=_GRIS)
    fila += 3

    ws.merge_cells(f"B{fila}:F{fila}")
    ws[f"B{fila}"] = f"Total a reembolsar: {_l(total_r)}"
    ws[f"B{fila}"].font = Font(bold=True, size=11, color=_AZUL)
    fila += 4

    for col, txt in (("C", "Solicitado por"), ("E", "Autorizado por")):
        ws[f"{col}{fila}"] = "_______________________"
        ws[f"{col}{fila+1}"] = txt
        ws[f"{col}{fila+1}"].font = Font(size=9)

    ruta = Path(ruta_salida) if ruta_salida else Path(REPORTES_DIR) / (
        f"memo_{(emp or 'empleado').replace(' ', '_')}_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(ruta))
    return str(ruta)


def generar_tabla(solicitudes, ruta_salida=None):
    """Tabla/planilla de liquidacion con el detalle completo de cada factura."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Liquidacion"
    cols = ["Folio", "Empleado", "Depto/Centro", "Zona", "Categoria", "RTN",
            "CAI", "Fecha", "Negocio", "Monto", "Limite", "Reembolsable",
            "Estado", "Observaciones"]
    for i, c in enumerate(cols):
        cel = ws.cell(row=1, column=1 + i, value=c)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=_AZUL)
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cel.border = _borde
    ws.freeze_panes = "A2"

    for r, s in enumerate(solicitudes, start=2):
        vals = [s.get("folio"), s.get("empleado"), s.get("departamento"),
                s.get("zona"), (s.get("categoria") or "").title(), s.get("rtn"),
                s.get("cai"), s.get("fecha_compra"), s.get("negocio"),
                float(s.get("monto") or 0), float(s.get("limite_aplicado") or 0),
                float(s.get("monto_reembolsable") or 0), s.get("estado"),
                s.get("observaciones") or s.get("motivos") or ""]
        for i, v in enumerate(vals):
            cel = ws.cell(row=r, column=1 + i, value=v)
            cel.border = _borde
            if i in (9, 10, 11):
                cel.number_format = '"L" #,##0.00'
    anchos = [16, 24, 16, 14, 16, 16, 24, 12, 22, 14, 14, 14, 14, 30]
    for i, w in enumerate(anchos):
        ws.column_dimensions[get_column_letter(1 + i)].width = w

    ruta = Path(ruta_salida) if ruta_salida else Path(REPORTES_DIR) / (
        f"liquidacion_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(ruta))
    return str(ruta)
